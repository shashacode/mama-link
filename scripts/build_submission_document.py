"""Build the submission-ready MAMA-Link multi-agent solution document."""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "MAMA-Link_Multi-Agent_Submission_Documentation.docx"
ASSET_DIR = ROOT / ".local" / "submission-document-assets"
FLOW_IMAGE = ASSET_DIR / "multi-agent-flow.png"

NAVY = "183A4A"
TEAL = "4D7773"
PALE_BLUE = "EAF1F4"
PALE_GREEN = "EDF3EF"
PALE_GOLD = "F6F0E5"
MID_GRAY = "D9D9D9"
TEXT_GRAY = "4E5962"
BLACK = "000000"
WHITE = "FFFFFF"


def set_run_font(run, name="Aptos", size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=115, bottom=105, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, text, header=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.clear()
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=header, color=WHITE if header else BLACK)


def prevent_row_split(row):
    """Keep each table row together when Word paginates the document."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for i, (header, width) in enumerate(zip(headers, widths)):
        set_cell_width(table.rows[0].cells[i], width)
        set_cell_shading(table.rows[0].cells[i], header_fill)
        format_cell(table.rows[0].cells[i], header, header=True, size=8.6)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        fill = PALE_BLUE if row_index % 2 else WHITE
        for i, (value, width) in enumerate(zip(values, widths)):
            set_cell_width(cells[i], width)
            set_cell_shading(cells[i], fill)
            format_cell(cells[i], value, size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, result, end])
    set_run_font(run, size=8.5, color=TEXT_GRAY)


def add_body(doc, text, bold_lead=None, after=6, keep=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items, level=0, size=10.2):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        set_run_font(p.add_run(item), size=size)


def add_numbered(doc, items, size=10.2):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        set_run_font(p.add_run(f"{index}.  {item}"), size=size)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(text), size=8.5, italic=True, color=TEXT_GRAY)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def rounded_box(draw, xy, fill, outline, title, subtitle, title_font, body_font):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    center = (x1 + x2) // 2
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text((center - (title_box[2] - title_box[0]) / 2, y1 + 22), title, fill=(0, 0, 0), font=title_font)
    lines = subtitle.split("\n")
    y = y1 + 68
    for line in lines:
        box = draw.textbbox((0, 0), line, font=body_font)
        draw.text((center - (box[2] - box[0]) / 2, y), line, fill=(48, 59, 65), font=body_font)
        y += 27


def arrow(draw, start, end, color=(77, 119, 115), width=6):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    if abs(end[0] - start[0]) > abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 11), (x2 - 18 * direction, y2 + 11)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x2, y2), (x2 - 11, y2 - 18 * direction), (x2 + 11, y2 - 18 * direction)]
    draw.polygon(points, fill=color)


def double_arrow(draw, start, end, color=(24, 58, 74), width=5):
    """Draw a two-way request/response connector."""
    arrow(draw, start, end, color=color, width=width)
    arrow(draw, end, start, color=color, width=width)


def create_flow_image():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 1100), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path("C:/Windows/Fonts")
    regular_path = font_dir / "arial.ttf"
    bold_path = font_dir / "arialbd.ttf"
    title_font = ImageFont.truetype(str(bold_path), 30)
    body_font = ImageFont.truetype(str(regular_path), 23)
    top_font = ImageFont.truetype(str(bold_path), 34)

    draw.text((900, 35), "MAMA-Link controlled multi-agent workflow", fill=(0, 0, 0), font=top_font, anchor="ma")

    # The left-hand control path shows who owns orchestration and the final output.
    rounded_box(draw, (60, 90, 600, 215), "#F6F0E5", "#C89B5B",
                "Web application or CLI", "Structured fictional case or\nconsented question", title_font, body_font)
    arrow(draw, (330, 215), (330, 270), color=(24, 58, 74))
    rounded_box(draw, (60, 270, 600, 425), "#EAF1F4", "#183A4A",
                "MAMA-Link orchestrator", "Invokes every stage with bounded inputs\nvalidates and aggregates every response", title_font, body_font)

    # A control bus makes the request/response relationship with every agent explicit.
    draw.line([(600, 347), (700, 347)], fill=(24, 58, 74), width=6)
    draw.line([(700, 152), (700, 922)], fill=(24, 58, 74), width=6)
    draw.text((1240, 64), "Numbered execution order; no direct agent-to-agent calls",
              fill=(78, 89, 98), font=body_font, anchor="ma")

    boxes = [
        ((800, 90, 1740, 215), "1  Intake Agent", "case ID -> structured summary | Tool  intake_summary"),
        ((800, 244, 1740, 369), "2  Triage Agent", "case ID -> risk classification | Tool  screen_case"),
        ((800, 398, 1740, 523), "3  Knowledge Agent", "screening topics -> cited education | Foundry File Search"),
        ((800, 552, 1740, 677), "4  Referral Agent", "case ID -> capable facility IDs | Tool  referral_plan"),
        ((800, 706, 1740, 831), "5  Transport Support Agent", "case ID -> support resource IDs | Tool  support_plan"),
        ((800, 860, 1740, 985), "6  Follow Up Agent", "case ID -> human review state | Tool  follow_up_plan"),
    ]
    for xy, title, subtitle in boxes:
        rounded_box(draw, xy, "#EDF3EF", "#4D7773", title, subtitle, title_font, body_font)
    for y in (152, 306, 460, 614, 768, 922):
        double_arrow(draw, (700, y), (800, y))

    arrow(draw, (330, 425), (330, 745), color=(24, 58, 74))
    draw.text((350, 585), "After all six responses\npass validation", fill=(78, 89, 98),
              font=body_font, anchor="lm")
    rounded_box(draw, (60, 745, 600, 900), "#EAF1F4", "#183A4A",
                "Validated workflow result", "Risk, grounded information, care options\nand review status", title_font, body_font)
    draw.text((900, 1045), "Two-way arrows show a bounded orchestrator request and validated agent response. Follow Up returns the final stage state; the orchestrator produces the result.",
              fill=(78, 89, 98), font=body_font, anchor="ma")
    image.save(FLOW_IMAGE, quality=95)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    for name, size, before, after in (("Heading 1", 16, 12, 7), ("Heading 2", 12.5, 10, 5), ("Heading 3", 10.8, 8, 4)):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for current_section in doc.sections:
        header = current_section.header
        header.distance = Cm(0.7)
        hp = header.paragraphs[0]
        hp.text = "MAMA Link Multi Agent Submission"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(hp.runs[0], size=8.5, color=TEXT_GRAY)
        footer = current_section.footer
        footer.distance = Cm(0.7)
        fp = footer.paragraphs[0]
        prefix = fp.add_run("Final activity submission   |   24 September 2026   |   ")
        set_run_font(prefix, size=8.5, color=TEXT_GRAY)
        add_page_number(fp)


def add_title_page(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(90)
    p.add_run("MAMA Link Multi Agent Maternal Health Coordination Solution")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    set_run_font(p.add_run("Submission documentation for the final activity Design and deliver a multi agent solution"), size=15, color=TEAL)

    for label, value in (
        ("Platform", "Microsoft Foundry Agent Service and FastAPI"),
        ("Solution stage", "Working hackathon MVP with live cloud agent verification"),
        ("Verification date", "24 September 2026"),
        ("Intended setting", "Maternal health coordination in underserved Nigerian communities"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        set_run_font(p.add_run(label + "  "), size=10.5, bold=True)
        set_run_font(p.add_run(value), size=10.5, color=TEXT_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("Submission statement"), size=11, bold=True)
    add_body(doc, "MAMA-Link coordinates bounded specialist agents to structure a maternal health case, classify synthetic risk, retrieve source-based education, match capable facilities and support resources, and prepare human follow-up. The system is a software demonstration. It does not diagnose, prescribe, confirm safety, dispatch emergency services, or replace qualified care.", after=10)
    add_body(doc, "All demonstrations use fictional data. Clinical content, screening rules, translations, facilities, support organisations and production operations require independent approval before real-world use.", bold_lead="All demonstrations use fictional data.")
    page_break(doc)


def build_document():
    create_flow_image()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "MAMA Link Multi Agent Maternal Health Coordination Solution"
    doc.core_properties.subject = "Final activity submission documentation"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "MAMA-Link, Microsoft Foundry, multi-agent, maternal health"
    add_title_page(doc)

    doc.add_heading("Document Purpose", level=1)
    add_body(doc, "This document explains the implemented MAMA-Link solution and the plan required to make it reliable, measurable and trustworthy. It covers every requested element of the final activity: the business problem, users, specialist agents, tools and knowledge sources, multi-agent rationale, workflow, observability, evaluation, governance, deployment and improvement lifecycle.")
    add_body(doc, "The main conclusion is that the architecture works as a controlled synthetic demonstration. On 24 September 2026, the patient-facing chat agent and the complete six-role Foundry workflow passed live tests against the existing gpt-4.1-mini deployment. Production use remains blocked by clinical validation, authenticated hosting, approved operational partners and independent evaluation.")

    doc.add_heading("Current Verification Snapshot", level=2)
    add_table(doc,
              ["Evidence", "Verified result"],
              [
                  ("Automated regression suite", "45 tests passed"),
                  ("Synthetic benchmark", "19 of 24 classifications matched; five expected-normal cases remained unassessed"),
                  ("Emergency regression set", "7 of 7 expected emergency cases detected"),
                  ("Referral capability check", "19 of 19 expected nonroutine cases matched required simulated capabilities"),
                  ("Live Foundry workflow", "Intake, triage, knowledge, referral, transport support and follow-up all passed for MAT-005"),
                  ("Patient-facing chat", "mama-link-chat v2 returned a consented educational answer through the localhost API"),
              ], [2.05, 4.65], font_size=9.1)

    doc.add_heading("Contents", level=1)
    contents = [
        "1 Business Problem and Outcome",
        "2 Intended Users",
        "3 Solution Scope and Boundaries",
        "4 Technical Architecture",
        "5 Specialist Agent Responsibilities",
        "6 Tools Data Sources and Knowledge",
        "7 Multi Agent Design Rationale",
        "8 Agent Workflow and Information Flow",
        "9 Demonstrated Workflow Result",
        "10 Production Readiness Plan",
        "11 Security Privacy and Safety",
        "12 Deployment and Operations",
        "13 Evaluation Evidence",
        "14 Limitations and Improvement Roadmap",
        "15 Submission Requirement Coverage",
        "References",
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(item), size=10.2)
    page_break(doc)

    doc.add_heading("1 Business Problem and Outcome", level=1)
    add_body(doc, "Pregnant and postpartum women in underserved communities may face several connected barriers at once. They may not recognise a warning sign, may not know whether the situation is urgent, may be directed to a facility without the required capability, or may lack transport and financial support. Caregivers and community health workers can face the same information and coordination gaps when helping someone else.")
    add_body(doc, "The business problem is therefore broader than symptom information. A useful service must coordinate risk recognition, trustworthy education, capability-based care matching, support discovery and follow-up without allowing a generative model to make an uncontrolled clinical or operational decision.")

    doc.add_heading("Business Outcome", level=2)
    add_body(doc, "MAMA-Link aims to shorten the path from a reported concern to an understandable and auditable next step. For a synthetic case, the solution produces a structured case summary, risk classification, source-based educational context, suitable facility options, relevant support options and a human review task. Consequential actions remain simulated and require consent and human authorisation.")

    doc.add_heading("Problem Constraints", level=2)
    add_bullets(doc, [
        "Medical uncertainty must remain visible. Missing data is unknown and an unmatched rule is unassessed, not normal.",
        "The nearest facility is not automatically appropriate. Matching must use required capabilities and available status.",
        "Health education must come from a controlled source set and clearly state its review status.",
        "No model may send a referral, book transport, notify a person or claim that an action occurred.",
        "The demo must work with fictional data and preserve account isolation, consent and auditability.",
    ])

    doc.add_heading("2 Intended Users", level=1)
    add_table(doc,
              ["User group", "Need supported by MAMA-Link", "Current interaction"],
              [
                  ("Pregnant and postpartum women", "Understand warning signs, ask educational questions, record symptoms and find appropriate care", "Personal account, symptom check, Ask Mama Link, learning area and history"),
                  ("Adolescent pregnant girls", "Accessible information and safeguarding-aware support pathways", "Same patient interface with adolescent support resources in the synthetic workflow"),
                  ("Family members and caregivers", "Help communicate a concern and prepare for escalation", "Can complete a structured check using fictional details in the MVP"),
                  ("Community health workers", "Structure observations and identify an appropriate referral pathway", "Web or CLI workflow for a selected synthetic case"),
                  ("Clinicians and facilities", "Receive a concise, traceable referral summary", "Simulated output only; no facility is contacted"),
                  ("NGOs and transport networks", "Match available support to location and need", "Fictional resource registry; no booking or contact"),
                  ("Programme administrators", "Review aggregate synthetic activity, test results and service health", "Admin-only case, evaluation and hotspot views"),
              ], [1.55, 3.05, 2.1], font_size=8.5)
    doc.add_heading("3 Solution Scope and Boundaries", level=1)
    add_body(doc, "The implemented MVP combines a local deterministic workflow with Azure Foundry prompt agents. The web application supports accounts, profiles, symptom checks, optional reviewed device-data imports, educational content, synthetic facility and support matching, longitudinal change flags and a consent-gated chat experience. The CLI demonstrates the complete multi-agent workflow.")

    doc.add_heading("Included in the MVP", level=2)
    add_bullets(doc, [
        "Fifteen versioned screening rules with critical precedence and an explicit unassessed outcome.",
        "Capability-based matching against fictional facility and support registries.",
        "Six sequential Foundry workflow roles plus a separate patient-facing chat agent.",
        "A review-pending WHO and UNFPA corpus accessed through Foundry File Search.",
        "Private local accounts, 12-hour server sessions, SQLite history and explicit cloud-context consent.",
        "Aggregate Application Insights metrics limited to operation name, status and duration.",
        "A 100-case synthetic hotspot demonstration and a 24-case labelled regression benchmark.",
    ])

    doc.add_heading("Excluded from the MVP", level=2)
    add_bullets(doc, [
        "Diagnosis, prescribing, dose selection, treatment instructions or declarations that a person is safe.",
        "Live ambulance dispatch, facility notification, transport booking, payment or caregiver messaging.",
        "Claims that fictional facilities are nearest, open or currently able to provide care.",
        "Native Apple Health or Android Health Connect synchronisation from the browser.",
        "Public hosting, production identity recovery, clinician-approved translations or clinical validation.",
    ])

    doc.add_heading("4 Technical Architecture", level=1)
    add_table(doc,
              ["Layer", "Implemented component", "Responsibility"],
              [
                  ("User experience", "Responsive HTML CSS and JavaScript", "Accounts, profiles, symptom entry, care search, history, education and consented chat"),
                  ("Application API", "FastAPI", "Authentication, validation, ownership checks, workflow endpoints and static delivery"),
                  ("Authoritative workflow", "Python rules and deterministic tools", "Screening, facility matching, support matching, provenance and follow-up state"),
                  ("Agent runtime", "Microsoft Foundry Agent Service", "Versioned prompt agents, function calling, File Search and response generation"),
                  ("Model", "gpt-4.1-mini deployment", "General reasoning under role-specific instructions; not a specialist medical model"),
                  ("Knowledge", "Foundry vector store", "Review-pending WHO and UNFPA educational corpus with source attribution"),
                  ("Persistence", "SQLite", "Local accounts, profiles, imports and assessment history"),
                  ("Monitoring", "Application Insights and Log Analytics", "Allow-listed aggregate operation count, status and duration"),
              ], [1.25, 2.15, 3.3], font_size=8.5)

    doc.add_heading("Architecture Control Principle", level=2)
    add_body(doc, "The model is not the system of record. The orchestrator supplies a bounded case identifier, executes a role-specific read-only tool and compares the agent's answer with the deterministic local result. A missing tool call, wrong case, extra argument, altered classification or changed identifier causes the run to fail. This design keeps screening and matching behaviour inspectable and repeatable.")
    doc.add_heading("5 Specialist Agent Responsibilities", level=1)
    add_body(doc, "The implemented design separates the workflow into six specialist roles. Ask Mama Link is a seventh prompt agent used only for patient-facing education. Agent versions below are the live versions verified on 24 September 2026.")
    add_table(doc,
              ["Agent", "Distinct responsibility", "Tool or source", "Validated output and boundary"],
              [
                  ("Intake v2", "Convert the selected synthetic case into a minimal structured summary", "intake_summary", "Case ID, pregnancy stage, location, reported symptoms and available reading names. Identity is excluded."),
                  ("Triage v3", "Report the authoritative screening classification", "screen_case", "Exactly one classification. It cannot downgrade or override the deterministic result."),
                  ("Knowledge v1", "Retrieve general education related to the screening topics", "Foundry File Search", "Answer must show a File Search call, cite the indexed source and retain the pending-review status."),
                  ("Referral v3", "Return facilities whose simulated capabilities meet the case requirements", "referral_plan", "Exact classification and ordered facility IDs. It cannot claim that a referral was sent."),
                  ("Transport Support v2", "Return matching fictional support resources", "support_plan", "Exact resource IDs and contacted false. It cannot contact or book a provider."),
                  ("Follow Up v2", "Prepare the required human review state", "follow_up_plan", "Review flag and task, with scheduled false and notification_sent false."),
                  ("Ask Mama Link Chat v2", "Answer general pregnancy and postpartum questions in plain language", "Supplied CDC education passages and consented profile context", "Educational response only. Emergency and medicine triggers use deterministic local guidance before any cloud call."),
              ], [1.1, 1.85, 1.25, 2.5], font_size=8.1)

    doc.add_heading("Agent Contract", level=2)
    add_body(doc, "The five tool-backed agents share contract version 2026-09-six-role-v1. The local manifest records each agent name and version, the chat deployment name and the project endpoint. The application refuses to treat chat as configured when the recorded model differs from the environment. This guard prevented the earlier stale gpt-4.1 reference from being presented as a working gpt-4.1-mini agent.")
    doc.add_heading("6 Tools Data Sources and Knowledge", level=1)
    doc.add_heading("Role Specific Tools", level=2)
    add_table(doc,
              ["Tool", "Input", "Authoritative function", "Side effects"],
              [
                  ("intake_summary", "case_id", "Returns a minimal structured case summary", "None"),
                  ("screen_case", "case_id", "Runs versioned screening rules", "None"),
                  ("referral_plan", "case_id", "Runs the workflow and returns capable simulated facilities", "None"),
                  ("support_plan", "case_id", "Returns matching fictional support resources", "None"),
                  ("follow_up_plan", "case_id", "Returns a human review task and unsent status flags", "None"),
                  ("file_search", "general education query", "Retrieves from the indexed public-source corpus", "None"),
              ], [1.3, 1.05, 3.1, 1.25], font_size=8.5)

    doc.add_heading("Data and Knowledge Sources", level=2)
    add_bullets(doc, [
        "Runtime maternal dataset: 100 synthetic cases. The first 24 retain benchmark labels for regression evaluation.",
        "Screening rules: version 0.3-hackathon with 15 deterministic rules and critical-over-warning precedence.",
        "Facility data: fictional offline records plus a link to the Nigeria Health Facility Registry for verification. Live capability and availability are not inferred.",
        "Support data: fictional organisations with service type, coverage and simulated availability.",
        "Knowledge corpus: review-pending summaries from WHO maternal mortality guidance, a WHO maternal and newborn counselling handbook, and UNFPA obstetric fistula material.",
        "Chat education: short CDC passages covering medicines in pregnancy, urgent maternal warning signs, postpartum concerns and pregnancy changes.",
    ])

    doc.add_heading("Knowledge Governance", level=2)
    add_body(doc, "Every knowledge item must carry its source URL, retrieval date, content owner, clinical reviewer, jurisdiction review, approval date, next review date and permitted use. The current corpus deliberately marks the reviewer fields as unassigned and permits hackathon retrieval only. An absent retrieval result cannot be interpreted as evidence of safety.")

    doc.add_heading("7 Multi Agent Design Rationale", level=1)
    add_body(doc, "A single general agent would have to interpret symptoms, retrieve health information, select facilities, match support and manage follow-up in one prompt. That would make permissions broad, errors difficult to locate and output validation weak. MAMA-Link separates these responsibilities so each role has a narrow instruction set, input schema, tool and acceptance condition.")
    add_table(doc,
              ["Design need", "Benefit of the multi-agent approach", "Control used in MAMA-Link"],
              [
                  ("Clear accountability", "A failure can be assigned to one role instead of an opaque combined answer", "Per-role result and tool-call log"),
                  ("Least privilege", "Each agent receives only the capability it needs", "One read-only tool per workflow agent"),
                  ("Consistent decisions", "Generative wording cannot change the authoritative classification or identifiers", "Exact output comparison with deterministic results"),
                  ("Grounded education", "Retrieval quality can be evaluated separately from triage and referral logic", "Dedicated File Search agent and source requirement"),
                  ("Maintainability", "A role, prompt or model can be versioned without redesigning every responsibility", "Stable agent references and explicit contract version"),
                  ("Failure containment", "A failed knowledge or chat call does not silently become a successful clinical decision", "Explicit unavailable state and local safety paths"),
              ], [1.35, 2.85, 2.45], font_size=8.45)
    add_body(doc, "The tradeoff is additional orchestration, latency, version management and cloud cost. The design accepts that overhead because safety-critical boundaries, observability and testability matter more than minimizing the number of agents.")
    doc.add_heading("8 Agent Workflow and Information Flow", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    flow_shape = p.add_run().add_picture(str(FLOW_IMAGE), width=Inches(6.65))
    flow_shape._inline.docPr.set("title", "MAMA-Link controlled multi-agent workflow")
    flow_shape._inline.docPr.set(
        "descr",
        "The web application or CLI sends a bounded case to the MAMA-Link orchestrator. Two-way arrows connect "
        "the orchestrator to six numbered agents, showing a bounded request and validated response for intake, "
        "triage, knowledge retrieval, referral, transport support and follow-up. After the Follow Up Agent returns "
        "the final review state, the orchestrator validates and aggregates all stage outputs into the workflow result.",
    )
    add_caption(doc, "Figure 1  Controlled sequence and bounded information flow")

    doc.add_heading("Sequence of Interactions", level=2)
    add_numbered(doc, [
        "The web application or CLI selects a fictional case. The orchestrator verifies the case ID and runs the local deterministic workflow.",
        "The orchestrator invokes the Intake Agent with the bounded case ID. The agent calls intake_summary and returns only the approved structured summary to the orchestrator.",
        "After validating the intake response, the orchestrator invokes the Triage Agent. The agent calls screen_case and returns the authoritative classification.",
        "The orchestrator derives general screening topics from the local workflow. The Knowledge Agent uses File Search and returns cited educational context.",
        "The orchestrator invokes the Referral Agent, which calls referral_plan and returns the ordered facility IDs that meet every required simulated capability.",
        "The orchestrator invokes the Transport Support Agent, which calls support_plan and returns matching resource IDs without contacting an organisation.",
        "The orchestrator invokes the Follow Up Agent, which calls follow_up_plan and returns the human review task with all action flags unsent and unscheduled.",
        "The Follow Up Agent does not create the final result. It returns the last stage state to the orchestrator, which validates it, aggregates all six outputs and returns a traceable workflow object to the demo process.",
    ], size=9.8)

    doc.add_heading("Information Passed Between Stages", level=2)
    add_body(doc, "Agents do not send unrestricted messages directly to one another. The orchestrator mediates the sequence, invokes every agent and validates each response before moving to the next stage. It passes the selected case ID to each tool-backed role and passes a general education query to the knowledge role. The query is derived from the local screening reasons rather than from another agent's unvalidated prose. The Follow Up Agent returns only a human review state; the orchestrator combines that state with the other validated outputs and the tool-call status list to produce the final result.")

    doc.add_heading("Patient Facing Chat Path", level=2)
    add_body(doc, "Ask Mama Link is intentionally separate from the six-role diagnostic simulation. After explicit consent, the API sends the user's question, a limited pregnancy summary, up to three recent assessments, a bounded conversation and supplied educational passages to mama-link-chat v2. Names, usernames and passwords are excluded from the automatically constructed context. Emergency and medicine questions take deterministic local paths before a cloud request.")
    doc.add_heading("9 Demonstrated Workflow Result", level=1)
    add_body(doc, "The live smoke test used synthetic case MAT-005, a third-trimester scenario in Idanre, Ondo, reporting no progress in labour, reduced fetal movement, prolonged or obstructed labour and severe exhaustion. No real person, facility or support provider was involved.")
    add_table(doc,
              ["Stage", "Live validated result"],
              [
                  ("Intake", "Structured MAT-005 summary with pregnancy stage, location, reported symptoms and available reading names"),
                  ("Triage", "critical"),
                  ("Knowledge", "File Search answer with WHO and UNFPA source URLs and pending clinical review status"),
                  ("Referral", "FAC-009, FAC-002 and FAC-004"),
                  ("Transport support", "RES-001, RES-002, RES-006 and RES-007; contacted false"),
                  ("Follow-up", "Human emergency-pathway review required; scheduled false; notification_sent false"),
                  ("Local workflow", "Emergency guidance, capability requirements and referral awaiting consent"),
              ], [1.55, 5.15], font_size=9.0)
    add_body(doc, "The test passed because every role called its required tool and returned the exact authoritative result. The knowledge role also showed a File Search call. The workflow would reject a downgraded classification, changed facility list, wrong case ID, missing tool call or claimed external action.")

    doc.add_heading("Final Output Delivered", level=2)
    add_body(doc, "For the business process, the final output is an auditable case package containing structured intake, risk status, source-based education, capable facility options, support options, follow-up state and the deterministic workflow trace. For a patient-facing chat question, the final output is a clearly labelled educational answer or an explicit unavailable message. Neither path claims that an external referral, notification or appointment occurred.")

    doc.add_heading("10 Production Readiness Plan", level=1)
    doc.add_heading("Observability Strategy", level=2)
    add_body(doc, "The current monitor records only allow-listed aggregate fields: operation, status and duration. Automatic FastAPI, request and Azure SDK capture is disabled so symptoms, prompts, notes, case IDs, session IDs, routes and tool payloads are not collected by default. Production observability should preserve that minimisation while adding role-level health signals.")
    add_table(doc,
              ["Signal", "Collection plan", "How it helps"],
              [
                  ("Operation count and status", "Count by workflow operation and success or failure", "Detects sudden failures and volume changes without patient content"),
                  ("End-to-end and role latency", "Histogram for API, each agent and each tool", "Locates slow roles, model throttling or downstream delays"),
                  ("Tool-call compliance", "Count missing, malformed, repeated and rejected calls by agent version", "Shows instruction drift or incompatible model behaviour"),
                  ("Output validation failures", "Count exact-match and schema rejections by role", "Identifies unsafe divergence before it reaches a user"),
                  ("Knowledge retrieval health", "File Search call rate, no-result rate and citation presence", "Detects indexing failures and corpus coverage gaps"),
                  ("Chat availability", "Configured, unavailable and timeout outcomes by version", "Separates deployment mismatch from model or identity failure"),
                  ("Consent and simulated referral state", "Aggregate accepted, declined and awaiting counts", "Confirms that consequential workflows remain consent-gated"),
                  ("Infrastructure health", "API availability, error rate, CPU, memory, database and Azure dependency status", "Supports capacity planning and incident response"),
              ], [1.45, 2.7, 2.55], font_size=8.25)

    doc.add_heading("Trace Design", level=3)
    add_body(doc, "A production trace should use a generated correlation ID and nested spans for orchestration, each agent invocation, each tool execution and validation. Tags may include agent role, agent version, contract version, tool name, status and duration. The trace must exclude prompts, symptoms, identifiers and raw outputs unless a separately approved, access-controlled diagnostic process is established.")
    doc.add_heading("Evaluation Strategy", level=2)
    doc.add_heading("Datasets and Testing Approach", level=3)
    add_bullets(doc, [
        "Regression dataset: the 24 labelled synthetic cases used to measure classification and referral capability matching.",
        "Workflow dataset: the 100 synthetic runtime cases used to exercise geographic and operational paths without representing prevalence.",
        "Agent contract tests: cases that verify tool use, exact outputs, wrong-case rejection, no side effects and temporary conversation deletion.",
        "Chat safety set: emergency wording, medicine questions, consent refusal, provider failure, prompt injection and context isolation.",
        "Held-out pre-production set: independently authored cases that were not used to design the screening rules.",
        "Clinician review set: approved scenarios and expected actions created and signed off by qualified maternal-health reviewers.",
    ])

    doc.add_heading("Evaluation Criteria", level=3)
    add_table(doc,
              ["Criterion", "Measure", "Proposed release gate"],
              [
                  ("Task adherence", "Required role output and no prohibited claims", "100 percent on mandatory safety scenarios"),
                  ("Tool usage", "Correct tool selected, called once with valid case ID and no extra arguments", "100 percent for tool-backed agents"),
                  ("Authoritative consistency", "Agent output exactly matches deterministic output", "100 percent for structured workflow fields"),
                  ("Groundedness", "Knowledge claims supported by retrieved approved content", "No unsupported clinical claim in reviewed evaluation set"),
                  ("Relevance and coherence", "Human and Foundry evaluator scores for educational responses", "Threshold approved after baseline and clinician review"),
                  ("Safety", "No diagnosis, prescription, false reassurance, dispatch claim or privacy leak", "Zero critical safety violations"),
                  ("Emergency sensitivity", "Expected emergency cases escalated", "100 percent on the approved emergency regression set"),
                  ("Reliability", "Availability, timeout rate and p95 latency", ">=99.5 percent pilot availability; latency targets set after representative load test"),
              ], [1.35, 3.0, 2.35], font_size=8.15)

    doc.add_heading("Development Lifecycle Integration", level=3)
    add_numbered(doc, [
        "Every change runs local unit, API, workflow and agent-boundary tests in continuous integration.",
        "Prompt, tool schema, rules or model changes create a new version and rerun the fixed regression suite.",
        "A staging gate runs held-out quality, groundedness, safety, tool-use and adversarial evaluations.",
        "A deliberate live smoke test confirms the exact Foundry agent version and deployment before promotion.",
        "Production samples are evaluated only under an approved privacy process. Aggregate monitoring continues continuously, and scheduled test-set evaluation checks for drift.",
        "Failures are reviewed by category, converted into regression cases and retested before a new version is released.",
    ], size=9.6)

    doc.add_heading("Governance and Reliability", level=2)
    add_table(doc,
              ["Control", "Current implementation", "Production requirement"],
              [
                  ("Version control", "Manifest records stable agent references and contract version", "Approval and rollback record for every promoted version"),
                  ("Least privilege", "One read-only case tool per workflow agent", "Managed identities and role assignments restricted by environment"),
                  ("Structured validation", "Strict schemas and exact authoritative-output comparison", "Schema compatibility tests and rejection alerts"),
                  ("Human authority", "All contact, scheduling and notification flags remain false", "Named human approvers and auditable confirmation for every consequential action"),
                  ("Failure behaviour", "Unassessed for unknown risk; unavailable for failed chat", "Documented degraded modes, escalation path and status communication"),
                  ("Knowledge control", "Restricted corpus with visible pending-review status", "Clinical owner, approval date, jurisdiction review and expiry workflow"),
                  ("Model compatibility", "Environment, manifest and live deployment names must agree", "Pre-deployment compatibility test and automated rollback"),
                  ("Change management", "45 automated tests and targeted cloud smoke tests", "Peer review, clinical sign-off and staged release evidence"),
              ], [1.25, 2.7, 2.75], font_size=8.15)
    doc.add_heading("11 Security Privacy and Safety", level=1)
    add_body(doc, "MAMA-Link applies data minimisation and fail-closed behaviour to the prototype. Passwords are salted and hashed with scrypt. Server-side sessions expire after 12 hours and are revoked at sign-out. Ownership checks prevent one account from reading or using another account's assessments or imports. Registration does not create an administrator.")
    add_table(doc,
              ["Risk", "Implemented mitigation", "Remaining production action"],
              [
                  ("Prompt injection", "Case text and tool output are declared untrusted; agents receive only bounded tools", "Add systematic jailbreak and indirect-injection evaluation"),
                  ("Tool misuse", "Exact case ID, role and argument checks; all tools are read-only", "Use managed identities, network controls and formal tool risk classification"),
                  ("Cross-user disclosure", "Server-side sessions and ownership checks; names and usernames excluded from constructed chat context", "Production IAM, recovery, penetration testing and access review"),
                  ("Sensitive telemetry", "Only aggregate operation, status and duration are emitted", "Data protection impact assessment, retention policy and audited access"),
                  ("Unsafe health response", "Deterministic emergency and medicine paths; no diagnosis or prescribing instruction", "Clinician-approved content, red-team testing and human escalation policy"),
                  ("Cloud outage", "Explicit unavailable response; learning cards and local safety messages remain", "Availability objectives, retry policy, circuit breaker and incident runbook"),
                  ("Stale model reference", "Manifest chat_model must match the configured deployment", "Deployment inventory check in release automation"),
                  ("Stale knowledge", "Corpus includes review status and source URLs", "Assigned owners, review dates and automatic expiry"),
              ], [1.3, 2.85, 2.55], font_size=8.15)

    doc.add_heading("Consent and Data Handling", level=2)
    add_body(doc, "A user may read learning content without sending information to a model. Ask Mama Link requires an explicit context-sharing choice. The cloud payload may include the question, pregnancy stage, age, gestational timing, medicines, allergies, a bounded conversation and up to three recent assessments. Automatically constructed context excludes account name, username and password. The model request uses store false, and the application does not persist chat messages in the database.")

    doc.add_heading("12 Deployment and Operations", level=1)
    add_body(doc, "The current web application runs on localhost and uses a local SQLite database. The Foundry agents, model deployment, File Search vector store and monitoring workspace exist in Azure. This split is appropriate for a hackathon demonstration but not for public service.")

    doc.add_heading("Recommended Deployment Path", level=2)
    add_table(doc,
              ["Stage", "Deployment activity", "Exit evidence"],
              [
                  ("Development", "Local API, synthetic data, unit tests and explicit agent version creation", "All automated tests pass and no secret is committed"),
                  ("Staging", "HTTPS hosting, managed identity, isolated database and non-production Foundry project", "Security review, load test, held-out evaluation and rollback rehearsal"),
                  ("Controlled pilot", "Limited approved users, qualified clinical oversight and verified partner directory", "Pilot objectives met with no critical safety or privacy incident"),
                  ("Production", "Scaled service, support rota, backup, recovery, change approval and audit retention", "Operational, legal, clinical and data-protection approval"),
              ], [1.1, 3.45, 2.15], font_size=8.5)

    doc.add_heading("Operational Improvement Loop", level=2)
    add_numbered(doc, [
        "Monitor aggregate service health, role latency, validation failures and retrieval health.",
        "Review incidents and failed evaluations without exposing patient content to unauthorised staff.",
        "Classify the cause as data, prompt, tool, model, knowledge, infrastructure or policy failure.",
        "Add a synthetic regression case and update only the responsible component.",
        "Run local tests, held-out evaluation and an intentional live smoke test.",
        "Promote a new version through approval and retain the previous manifest for rollback.",
    ], size=9.7)

    doc.add_heading("13 Evaluation Evidence", level=1)
    add_table(doc,
              ["Evidence item", "Result", "Interpretation"],
              [
                  ("Automated suite", "45 of 45 passed", "Covers API security, account isolation, workflow rules, imports, chat safety and agent boundaries"),
                  ("Classification", "19 of 24 matched", "Five expected-normal fixtures remain unassessed because a clinically approved low-risk completion protocol is absent"),
                  ("Emergency detection", "7 of 7", "All known critical fixtures entered the emergency pathway; this is regression evidence, not clinical sensitivity"),
                  ("Referral capability", "19 of 19 expected nonroutine cases", "Matched simulated capabilities under the documented crosswalk"),
                  ("Live six-role run", "Passed on MAT-005", "Every agent used its required tool or File Search and returned a validated result"),
                  ("Live Ask Mama Link", "Passed through localhost API", "mama-link-chat v2 used gpt-4.1-mini after explicit consent"),
              ], [1.45, 1.55, 3.7], font_size=8.55)
    add_body(doc, "The rules were developed with the known fixtures, so the benchmark measures regression coverage rather than generalisation or clinical validity. The northern weighting of the 100-case map is synthetic and cannot be interpreted as prevalence or a prediction about any community.")
    doc.add_heading("14 Limitations and Improvement Roadmap", level=1)
    add_table(doc,
              ["Limitation", "Why it matters", "Next action"],
              [
                  ("No clinical validation", "Software correctness does not establish medical correctness", "Qualified Nigerian maternal-health reviewers approve rules, messages and evaluation cases"),
                  ("Expected-normal cases remain unassessed", "The system cannot yet support a low-risk completion claim", "Define a clinically approved minimum-data and exclusion protocol"),
                  ("Fictional facilities and support providers", "The demo cannot confirm real capability, availability or contact details", "Integrate the official registry and verified partners through server-side credentials"),
                  ("No real dispatch", "The system prepares but does not execute emergency coordination", "Design authorised human approval and acknowledgement workflows"),
                  ("General education model", "gpt-4.1-mini is not a medically trained or clinically validated model", "Compare approved models and prompts on the clinician-reviewed evaluation set"),
                  ("Review-pending corpus", "Source quality alone does not approve local clinical use", "Assign a content owner, jurisdiction reviewer and review schedule"),
                  ("Local authentication and storage", "Local sessions and SQLite are unsuitable for a public health service", "Deploy managed identity, production database, encryption, recovery and audit controls"),
                  ("Partial language and device support", "Translations and browser imports do not prove accessible nationwide operation", "Professional language review, usability research and native mobile permission flows"),
              ], [1.45, 2.5, 2.75], font_size=8.25)

    doc.add_heading("Prioritised Roadmap", level=2)
    add_numbered(doc, [
        "Clinical and jurisdiction review of the rules, knowledge corpus, user messages and low-risk completion criteria.",
        "Independent held-out and adversarial evaluation with approved release thresholds.",
        "Authenticated staging deployment with managed identity, secure secrets, production database and privacy assessment.",
        "Verified health facility and support partner integrations with human approval for every consequential action.",
        "Controlled pilot, continuous monitoring, incident response drills and version rollback testing.",
    ], size=9.8)

    doc.add_heading("15 Submission Requirement Coverage", level=1)
    add_table(doc,
              ["Activity requirement", "Where it is addressed"],
              [
                  ("Business problem", "Section 1 defines the coordination and access problem and the intended outcome"),
                  ("Intended users", "Section 2 identifies primary, assisting, clinical and programme users"),
                  ("At least two specialised agents", "Section 5 defines six workflow agents and a separate chat agent"),
                  ("Tools data sources and knowledge bases", "Section 6 lists every tool, dataset, registry and knowledge source"),
                  ("Why multi-agent is more effective", "Section 7 explains accountability, least privilege, validation and failure containment"),
                  ("Information flow between agents", "Section 8 provides the sequence, diagram and bounded payload description"),
                  ("Observability strategy", "Section 10 defines traces, metrics, privacy boundaries and diagnostic use"),
                  ("Evaluation strategy", "Section 10 defines datasets, criteria, release gates and lifecycle integration"),
                  ("Governance and reliability", "Sections 10 and 11 define versioning, validation, permissions, review and failure behaviour"),
                  ("Workflow and final output", "Sections 8 and 9 show the sequence, tool calls, live result and business output"),
                  ("Deployment monitoring evaluation and improvement", "Sections 10, 12 and 14 define staged deployment and the improvement loop"),
              ], [2.3, 4.4], font_size=8.65)

    doc.add_heading("Submission Conclusion", level=2)
    add_body(doc, "MAMA-Link demonstrates a practical multi-agent pattern for a sensitive domain: generative agents explain and coordinate, while deterministic tools and explicit validation retain authority over structured decisions. The current implementation satisfies the final activity as a working synthetic demonstration and includes a concrete path toward measurable, governed production readiness. It must remain outside clinical use until the listed clinical, operational, security and data-governance requirements are completed.")
    doc.add_heading("References", level=1)
    references = [
        ("World Health Organization  Maternal mortality", "https://www.who.int/news-room/fact-sheets/detail/maternal-mortality"),
        ("World Health Organization  Counselling for Maternal and Newborn Health Care", "https://iris.who.int/bitstream/handle/10665/44016/9789241547628_eng.pdf"),
        ("United Nations Population Fund  Obstetric fistula", "https://www.unfpa.org/obstetric-fistula"),
        ("Centers for Disease Control and Prevention  Urgent maternal warning signs", "https://www.cdc.gov/hearher/maternal-warning-signs/index.html"),
        ("Centers for Disease Control and Prevention  Medicine and pregnancy", "https://www.cdc.gov/medicine-and-pregnancy/about/index.html"),
        ("Microsoft Learn  Prompt agent quickstart", "https://learn.microsoft.com/en-us/azure/foundry/agents/quickstarts/prompt-agent"),
        ("Microsoft Learn  Function calling with Foundry agents", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/function-calling"),
        ("Microsoft Learn  Observability in generative AI", "https://learn.microsoft.com/en-us/azure/foundry/concepts/observability"),
        ("Nigeria Health Facility Registry", "https://hfr.fmohconnect.gov.ng/facilitieslist"),
        ("Apple Developer  Setting up HealthKit", "https://developer.apple.com/documentation/healthkit/setting-up-healthkit"),
        ("Android Developers  Health Connect get started", "https://developer.android.com/health-and-fitness/health-connect/get-started"),
    ]
    for index, (title, url) in enumerate(references, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"{index}. {title}. "), size=9.4)
        add_hyperlink(p, url, url)

    doc.add_heading("Project Evidence", level=2)
    add_bullets(doc, [
        "README.md and docs/foundry-setup.md for current runtime, agent versions and live verification status.",
        "docs/evaluation-latest.json for the 24-case regression results.",
        "tests for API, personal data, workflow and Foundry agent boundary coverage.",
        "knowledge/public-sources/maternal-health-public-sources.md for the current retrieval corpus and review record.",
    ], size=9.4)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
