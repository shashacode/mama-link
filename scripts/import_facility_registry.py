"""Import the Phase 1 Nigeria facility workbook into the runtime directory."""
import json
import re
import gzip
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Nigeria_Healthcare_Facility_Registry_Phase_1.xlsx"
OUTPUT = ROOT / "mama_link_dataset" / "maternal" / "data" / "facility_registry.json.gz"


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "").lstrip("\ufeff") for value in next(iterator)]
    return [dict(zip(headers, values)) for values in iterator if any(value is not None for value in values)]


def text(value):
    return str(value).strip() if value is not None else ""


def coordinate(value):
    cleaned = text(value).replace("`", "").strip(" ,:")
    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", cleaned):
        return None
    return float(cleaned)


def import_registry():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    all_rows = rows(workbook["All Facilities"])
    priority_rows = {text(row.get("facility_code")): row for row in rows(workbook["Maternal-Emergency Priority"])}
    facilities = []
    for row in all_rows:
        code = text(row.get("facility_code"))
        priority = priority_rows.get(code, {})
        longitude = row.get("longitude")
        latitude = row.get("latitude")
        facilities.append({
            "facility_id": f"HFR-{code}",
            "registry": "Nigeria Healthcare Facility Registry",
            "registry_code": code,
            "uid": text(row.get("uid")),
            "name": text(row.get("facility_name")),
            "state": text(row.get("state")),
            "lga": text(row.get("lga")),
            "ward": text(row.get("ward")),
            "ownership": text(row.get("ownership")),
            "facility_level": text(row.get("facility_level")),
            "operation_status": text(row.get("operation_status")),
            "registration_status": text(row.get("registration_status")),
            "license_status": text(row.get("license_status")),
            "referral_level_proxy": text(row.get("referral_level_proxy")),
            "research_priority": text(row.get("research_priority")),
            "priority_reason": text(row.get("priority_reason")),
            "latitude": coordinate(latitude),
            "longitude": coordinate(longitude),
            "street_address": text(priority.get("street_address")),
            "public_phone": text(priority.get("public_phone")),
            "public_email": text(priority.get("public_email")),
            "official_website": text(priority.get("official_website")),
            "contact_verification_status": text(priority.get("contact_verification_status")),
            "contact_verification_date": text(priority.get("contact_verification_date")),
            "capabilities": [],
            "registry_entry": True,
            "source_workbook": SOURCE.name,
            "source_status": "Phase 1 registry import; contact and service fields require verification"
        })
    payload = {
        "source": SOURCE.name,
        "source_sheet": "All Facilities plus Maternal-Emergency Priority contact queue",
        "record_count": len(facilities),
        "states": sorted({facility["state"] for facility in facilities}),
        "disclaimer": "Registry directory data. Do not infer clinical capability, live availability, or contact validity from this import.",
        "facilities": facilities
    }
    with gzip.open(OUTPUT, "wt", encoding="utf-8") as handle:
        json.dump(payload, handle, separators=(",", ":"))
    print(f"Imported {len(facilities)} registry facilities from {SOURCE.name}.")


if __name__ == "__main__":
    import_registry()